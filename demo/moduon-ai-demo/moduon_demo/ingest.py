"""수집 파이프라인의 공용 단계 — 터미널 시연(scenes.py)과 실행 프로그램(app/)이 함께 쓴다.

여기 있는 함수는 화면 출력을 하지 않는다. 원본 보관, 규칙 파싱, AI 출력 검증, staging 저장, 규칙 매칭,
이상 판정처럼 '코드가 하는 일'만 한다. AI 호출은 부르는 쪽(scenes.py, app/service.py)이 한다.
"""
import csv
import hashlib
import json
import re
from pathlib import Path

import pdfplumber
from openpyxl import load_workbook
from openpyxl.utils import get_column_letter

from .rules import anomaly as anomaly_rules
from .rules import verify
from .rules.conditions import JOIN_KEYS, PlanBook, UnknownCondition
from .rules.matching import Catalog
from .rules.text import compact, parse_count, parse_krw

MONTH, PREV_MONTH = "2026-10", "2026-09"
FIELD_UNIT = {"monthly_rental_fee": "KRW", "mandatory_months": "month", "registration_fee": "KRW"}
C_TEMPLATE = ["상품코드", "단말명", "출고가", "요금제코드", "공시지원금", "리베이트_번호이동", "리베이트_기기변경"]
PII = re.compile(r"(\d{6}-?[1-4]\d{6})|(01[016789]-?\d{3,4}-?\d{4})")
PLACEHOLDER = re.compile(r"\{(prev_value|new_value|change_pct)\}")
DEFAULT_TEMPLATE = "이전 값 {prev_value}에서 {new_value}로 바뀌었습니다({change_pct})."


def sha(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def won(v) -> str:
    return "-" if v is None else f"{v:,}"


def fix_particles(text: str) -> str:
    """'690,000원로' → '690,000원으로' — 숫자를 코드가 채운 뒤 받침에 맞게 조사를 고친다."""
    return re.sub(r"원(로|를|는)(?=\s|[.,)(]|$)", lambda m: "원" + {"로": "으로", "를": "을", "는": "은"}[m[1]], text)


def grade(checks) -> str:
    """등급은 모델이 말한 확신도가 아니라 검증 결과로만 정한다."""
    if any(c.ok is False for c in checks):
        return "low"
    return "high" if all(c.ok for c in checks) else "medium"


def save_raw(st, partner: str, kind: str, filename: str, data: bytes, injection_suspect: bool = False) -> int:
    with st.role("ingest_worker"):
        return st.exec("insert into raw_file(partner, kind, filename, sha256, injection_suspect) values (?,?,?,?,?)",
                       (partner, kind, filename, sha(data), int(injection_suspect)))


def add_mention(st, rid, partner, category, raw_name, model_code=None) -> int:
    with st.role("ingest_worker"):
        return st.exec("insert into staging_mention(raw_file_id, partner, category, raw_name, model_code) values (?,?,?,?,?)",
                       (rid, partner, category, raw_name, model_code))


def add_record(st, **r) -> int:
    cols = ["raw_file_id", "mention_id", "partner", "plan_id", "field_code", "condition_key", "value_int", "unit",
            "value_text", "loc", "evidence_text", "extractor", "grade", "checks", "state", "note"]
    r.setdefault("plan_id", "")
    r.setdefault("state", "validated")
    with st.role("ingest_worker"):
        return st.exec(f"insert into staging_record({', '.join(cols)}) values ({', '.join('?' * len(cols))})",
                       tuple(r.get(c) for c in cols))


def add_note(st, rid, partner, kind, text, mention_id=None):
    with st.role("ingest_worker"):
        st.exec("insert into staging_note(raw_file_id, partner, mention_id, kind, text) values (?,?,?,?,?)",
                (rid, partner, mention_id, kind, text))


# ───────────────────────────── 기존 확정 데이터(상품·요금제 마스터, alias, 지난달 확정값)
def load_masters(st, root: Path) -> dict:
    n = {"products": 0, "plans": 0, "aliases": 0, "prices": 0}
    with st.role("setup"):
        for r in csv.DictReader(open(root / "data/product_master.csv", encoding="utf-8")):
            st.exec("insert into canonical_product values (?,?,?,?,?,?,?,?)",
                    (r["id"], r["category"], r["vendor"], r["name"], r["model_code"], r["storage_gb"], r["color"],
                     r["variant"]))
            n["products"] += 1
        for r in csv.DictReader(open(root / "data/plan_master.csv", encoding="utf-8")):
            st.exec("insert into canonical_plan values (?,?,?,?)", (r["partner"], r["plan_id"], r["name"], int(r["monthly_fee"])))
            for a in r["aliases"].split("|"):
                st.exec("insert into canonical_plan_alias values (?,?,?)", (r["partner"], a, r["plan_id"]))
            n["plans"] += 1
        for r in csv.DictReader(open(root / "data/product_alias.csv", encoding="utf-8")):
            st.exec("insert into canonical_alias values (?,?,?,?)", (r["partner"], r["alias"], r["product_id"], "기존 확정"))
            n["aliases"] += 1
        for r in csv.DictReader(open(root / "data/price_2609.csv", encoding="utf-8")):
            st.exec("insert into canonical_price(partner, product_id, plan_id, field_code, condition_key, value_int, unit,"
                    " month, source, approved_by) values (?,?,?,?,?,?,?,?,?,?)",
                    (r["partner"], r["product_id"], r["plan_id"], r["field_code"], r["condition_key"],
                     int(r["value_int"]), r["unit"], PREV_MONTH, "기존 확정", "기존 확정"))
            n["prices"] += 1
    return n


# ───────────────────────────── 엑셀 정책표
def excel_grid(path_or_bytes, sample_rows: int = 10):
    """(시트, {(행, 열): 글자}, AI에게 보낼 텍스트, 열 목록). 병합 셀은 코드가 같은 글자로 채운다."""
    import io
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    ws = load_workbook(src, data_only=True).active
    fill = {}
    for rng in ws.merged_cells.ranges:
        v = ws.cell(rng.min_row, rng.min_col).value
        for r in range(rng.min_row, rng.max_row + 1):
            for c in range(rng.min_col, rng.max_col + 1):
                fill[(r, get_column_letter(c))] = v
    cols = [get_column_letter(c) for c in range(1, ws.max_column + 1)]
    grid, lines = {}, []
    for r in range(1, sample_rows + 1):
        cells = []
        for L in cols:
            v = fill.get((r, L), ws[f"{L}{r}"].value)
            grid[(r, L)] = "" if v is None else str(v)
            if v is not None:
                cells.append(f"{L}: {v}")
        if cells:
            lines.append(f"행 {r} | " + " | ".join(cells))
    return ws, grid, "\n".join(lines), cols


def excel_parse(st, rid: int, partner: str, category: str, ws, grid: dict, cols: list[str], final: dict,
                header_rows: list[int]) -> tuple[list, list]:
    """승인된 매핑으로 '규칙 파서'가 전체 행을 읽어 staging에 넣는다(AI 아님). (가격 열 목록, 행별 [상품명, 값…])"""
    prod_col = next(c for c, v in final.items() if v[0] == "product_name")
    model_col = next((c for c, v in final.items() if v[0] == "model_code"), None)
    memo_col = next((c for c, v in final.items() if v[0] == "memo"), None)
    price_cols = [(c, v) for c, v in final.items() if v[0] in verify.PRICE_FIELDS]
    heads = {L: verify.compose_header(grid, header_rows, L) for L in cols}
    for r in range(1, min(header_rows)):          # 열 제목 위의 안내문(단위·조건)은 조건 메모로 보관
        text = ws[f"A{r}"].value
        if text and str(text).startswith("※"):
            add_note(st, rid, partner, "sheet_note", str(text))
    rows = []
    for r in range(max(header_rows) + 1, ws.max_row + 1):
        name = ws[f"{prod_col}{r}"].value
        if not name:
            continue
        model = ws[f"{model_col}{r}"].value if model_col else None
        memo = ws[f"{memo_col}{r}"].value if memo_col else None
        source_row = f"행 {r}: " + " | ".join(f"{heads[L]}={ws[f'{L}{r}'].value}" for L in cols
                                                if ws[f"{L}{r}"].value is not None)
        mid = add_mention(st, rid, partner, category, str(name), model)
        values = []
        for c, (fcode, plan_id, cond, unit) in price_cols:
            raw_v = ws[f"{c}{r}"].value
            v = raw_v if isinstance(raw_v, int) else parse_krw(str(raw_v))
            value = None if v is None else v * verify.UNIT_MULT.get(unit, 1)
            add_record(st, raw_file_id=rid, mention_id=mid, partner=partner, plan_id=plan_id, field_code=fcode,
                       condition_key=cond, value_int=value, unit="KRW", value_text=str(raw_v), loc=f"{c}{r}",
                       evidence_text=source_row, extractor="rule", grade="rule", note=memo)
            values.append(value)
        rows.append([str(name), *values])
    return price_cols, rows


# ───────────────────────────── 카톡·문자 공지
def kakao_verify(text: str, ai: dict, plans: PlanBook, partner: str) -> list:
    """[(AI 레코드, 검사 목록, 코드가 해석한 값, 등급)]"""
    out = []
    for rec in ai["records"]:
        checks, got = verify.kakao_checks(text, rec, plans, partner)
        out.append((rec, checks, got, grade(checks)))
    return out


def kakao_score(staged: list, expected: list[dict]) -> tuple[int, int]:
    """(맞은 칸 수, 정답에 없는 칸 수). 같은 칸에 다른 값도 적었으면 틀림."""
    exp = {(compact(k["raw_name"]), k["plan_id"], k["condition_key"]): k["value"] for k in expected}
    got = {}
    for rec, _, g, _ in staged:
        got.setdefault((compact(rec["product_ref_raw"]), g["plan_id"], g["condition_key"]), set()).add(g["value"])
    return sum(1 for k, v in exp.items() if got.get(k) == {v}), sum(1 for k in got if k not in exp)


def kakao_store(st, rid: int, partner: str, text: str, staged: list, conditions: list[str]) -> int:
    """검증 결과와 함께 staging에 저장. 원문에 있는 조건 문장만 메모로 보관. (보관한 조건 수)"""
    mids = {}
    for rec, checks, got, g in staged:
        name = rec["product_ref_raw"]
        if compact(name) not in mids:
            mids[compact(name)] = add_mention(st, rid, partner, "telecom", name)
        add_record(st, raw_file_id=rid, mention_id=mids[compact(name)], partner=partner, plan_id=got["plan_id"] or "",
                   field_code="rebate", condition_key=got["condition_key"] or "base", value_int=got["value"], unit="KRW",
                   value_text=rec["value_text"], loc=rec["join_phrase"], evidence_text=rec["evidence_text"],
                   extractor="ai", grade=g, checks=json.dumps({c.name: c.ok for c in checks}, ensure_ascii=False))
    kept = 0
    for c in conditions:
        if verify.evidence_in_text(text, c).ok:
            add_note(st, rid, partner, "condition", c)
            kept += 1
    return kept


# ───────────────────────────── 양식이 고정된 전산(AI 없음)
def feed_api(st, path: Path, partner: str, plans: PlanBook) -> dict:
    body = json.loads(path.read_text(encoding="utf-8"))
    rid = save_raw(st, partner, "api", path.name, path.read_bytes())
    n = 0
    for it in body["items"]:
        mid = add_mention(st, rid, partner, "telecom", it["device_code"])
        recs = [("", "device_price", it["release_price"], it["device_code"])]
        recs += [(plans.resolve(partner, s["plan_code"]), "subsidy_amount", s["amount"],
                  f"{it['device_code']}/{s['plan_code']}") for s in it["support"]]
        for plan_id, fcode, v, loc in recs:
            add_record(st, raw_file_id=rid, mention_id=mid, partner=partner, plan_id=plan_id, field_code=fcode,
                       condition_key="base", value_int=v, unit="KRW", value_text=str(v), loc=loc, extractor="api",
                       grade="rule")
            n += 1
    return {"raw_file_id": rid, "request": body.get("request", ""), "devices": len(body["items"]), "values": n}


def feed_csv(st, path_or_rows, filename: str, partner: str, plans: PlanBook, data: bytes) -> dict:
    """승인된 양식(C_TEMPLATE)의 CSV. 열 제목이 다르면 읽지 않는다."""
    rows = path_or_rows
    if rows[0] != C_TEMPLATE:
        return {"raw_file_id": None, "template_ok": False, "devices": 0, "values": 0}
    rid = save_raw(st, partner, "csv", filename, data)
    n, mids = 0, {}
    for code, _, price, plan_code, sub, mnp, chg in rows[1:]:
        recs = []
        if code not in mids:
            mids[code] = add_mention(st, rid, partner, "telecom", code)
            recs.append(("", "device_price", "base", int(price), code))
        pid = plans.resolve(partner, plan_code)
        recs += [(pid, f, c, int(v), f"{code}/{plan_code}") for f, c, v in
                 [("subsidy_amount", "base", sub), ("rebate", "join=mnp", mnp), ("rebate", "join=chg", chg)]]
        for plan_id, fcode, cond, v, loc in recs:
            add_record(st, raw_file_id=rid, mention_id=mids[code], partner=partner, plan_id=plan_id, field_code=fcode,
                       condition_key=cond, value_int=v, unit="KRW", value_text=str(v), loc=loc, extractor="template",
                       grade="rule")
            n += 1
    return {"raw_file_id": rid, "template_ok": True, "devices": len(mids), "values": n}


# ───────────────────────────── PDF 안내문
def pdf_read(path_or_bytes):
    import io
    src = io.BytesIO(path_or_bytes) if isinstance(path_or_bytes, bytes) else path_or_bytes
    with pdfplumber.open(src) as pdf:
        pages = [p.extract_text() or "" for p in pdf.pages]
        tables = [p.extract_tables() for p in pdf.pages]
    return pages, tables, "\n".join(pages)


def pdf_verify(pages, tables, ai: dict) -> list:
    """[(행, 필드, 검사 목록, 등급)] — 근거·셀·숫자 3중 대조."""
    staged = []
    for row in ai["rows"]:
        pi = row["page"] - 1
        ptext = pages[pi] if 0 <= pi < len(pages) else ""
        ptables = tables[pi] if 0 <= pi < len(tables) else []
        for f in row["fields"]:
            unit = FIELD_UNIT.get(f["field"], "KRW")
            checks = [verify.evidence_match(ptext, f["evidence_text"], row["row_label"], f["value_text"]),
                      verify.cell_match(ptables, row["row_label"], f["col_header"], f["value_text"]),
                      verify.number_crosscheck(f["value_text"], f["value_number"], unit)]
            staged.append((row, f, checks, grade(checks)))
    return staged


def pdf_store(st, rid: int, partner: str, staged: list) -> int:
    mids, n = {}, 0
    for row, f, checks, g in staged:
        if f["field"] == "other":
            continue
        lbl = row["row_label"]
        if lbl not in mids:
            mids[lbl] = add_mention(st, rid, partner, "appliance", lbl)
        unit = FIELD_UNIT[f["field"]]
        value = parse_krw(f["value_text"]) if unit == "KRW" else parse_count(f["value_text"])
        add_record(st, raw_file_id=rid, mention_id=mids[lbl], partner=partner, field_code=f["field"],
                   condition_key="base", value_int=value, unit=unit, value_text=f["value_text"],
                   loc=json.dumps({"page": row["page"], "row": lbl, "col": f["col_header"]}, ensure_ascii=False),
                   evidence_text=f["evidence_text"], extractor="ai", grade=g,
                   checks=json.dumps({c.name: c.ok for c in checks}, ensure_ascii=False))
        n += 1
    return n


# ───────────────────────────── 상품명 매칭(규칙 단계)
def rule_match(st, mentions: list[dict]) -> tuple[list, list]:
    """규칙(M0 alias·모델코드, M1 정규화)으로 풀리면 바로 연결. (규칙 표 행, AI 판정이 필요한 [(mention, 후보)])"""
    cat = Catalog.from_db(st)
    rule_rows, need_ai = [], []
    with st.role("ingest_worker"):
        for m in mentions:
            raw, partner, category = m["raw_name"], m["partner"], m["category"]
            pid, path = cat.m0_alias(partner, raw), "M0 alias(전에 승인한 이름·연동 코드)"
            if not pid:
                ids = cat.m0_model_code(m["model_code"])
                pid, path = (ids[0], f"M0 모델코드 {m['model_code']} 유일 일치") if len(ids) == 1 else (None, None)
                if len(ids) > 1:
                    rule_rows.append([partner, raw, "M0 모델코드", f"{m['model_code']} → 후보 {len(ids)}개(유일하지 않음)"])
            if not pid:
                ids = cat.m1_norm(raw, category)
                pid, path = (ids[0], "M1 정규화 규칙 일치") if len(ids) == 1 else (None, None)
            if pid:
                st.exec("update staging_mention set product_id=?, match_state='matched_rule', match_path=? where id=?",
                        (pid, path, m["id"]))
                rule_rows.append([partner, raw, path, f"→ {cat.products[pid].name} (AI 호출 안 함)"])
            else:
                cands = cat.m2_candidates(raw, category, 5)
                need_ai.append((m, cands))
                rule_rows.append([partner, raw, "M2 유사도 후보", f"후보 {len(cands)}개 → AI 판정 필요"])
    return rule_rows, need_ai


COLOR_KO = {"black": "블랙", "silver": "실버", "white": "화이트"}


def match_block(cat: Catalog, mk: str, m: dict, cands: list) -> str:
    """AI에게 보낼 '이름 + 후보' 한 덩어리."""
    mc = m["model_code"] if m["model_code"] and m["model_code"] != "-" else "없음"
    lines = [f"[{mk}] 파트너: {m['partner']}", f'  원래 이름: "{m["raw_name"]}"', f"  파일의 모델코드: {mc}",
             "  후보(유사도 순):"]
    for j, (pid, _) in enumerate(cands, 1):
        p = cat.products[pid]
        attrs = " / ".join(x for x in [f"모델코드 {p.model_code}" if p.model_code else "",
                                        f"용량 {p.storage_gb}GB" if p.storage_gb else "",
                                        f"색상 {COLOR_KO.get(p.color, p.color)}" if p.color else "",
                                        f"옵션 {p.variant}" if p.variant else ""] if x)
        lines.append(f"    c{j}: {p.name}" + (f" ({attrs})" if attrs else ""))
    return "\n".join(lines)


# ───────────────────────────── 이상 판정(규칙) + 설명 틀 채우기
def detect(st, records: list[dict]) -> list:
    """지난달 확정값과 비교해 상태를 정하고 이상 이벤트를 저장한다. [(레코드, 이전 값, 이벤트, 상태)]"""
    out = []
    with st.role("ingest_worker"):
        for r in records:
            prev = st.one("select value_int from canonical_price where partner=? and product_id=? and plan_id=? "
                          "and field_code=? and condition_key=? and month=?",
                          (r["partner"], r["product_id"], r["plan_id"], r["field_code"], r["condition_key"], PREV_MONTH))
            prev_v = prev["value_int"] if prev else None
            if prev_v == r["value_int"]:
                st.exec("update staging_record set state='unchanged' where id=?", (r["id"],))
                out.append((r, prev_v, [], "unchanged"))
                continue
            events = anomaly_rules.check(r["field_code"], prev_v, r["value_int"]) if r["value_int"] is not None else \
                [{"rule_code": "UNPARSEABLE", "severity": "block", "reason": "값을 숫자로 읽을 수 없음"}]
            for e in events:
                st.exec("insert into staging_anomaly(record_id, rule_code, severity, reason, prev_value, new_value) "
                        "values (?,?,?,?,?,?)", (r["id"], e["rule_code"], e["severity"], e["reason"], prev_v, r["value_int"]))
            state = "blocked" if any(e["severity"] == "block" for e in events) else "pending_review"
            st.exec("update staging_record set state=? where id=?", (state, r["id"]))
            out.append((r, prev_v, events, state))
    return out


def fill_template(tmpl: str, prev_v, new_v) -> tuple[str, str | None]:
    """AI가 쓴 설명 틀에 코드가 숫자를 채운다. (완성 문장, 틀을 버린 이유 또는 None)

    허용된 자리표시자 3개 밖에 숫자나 다른 중괄호({memo} 등)가 있으면 틀을 버리고 기본 문구를 쓴다."""
    rest = PLACEHOLDER.sub("", tmpl)
    why = None
    if not tmpl.strip():
        why = "설명이 비어 있음"
    elif m := re.search(r"\{[^{}]*\}", rest):
        why = f"허용되지 않은 자리표시자 {m[0]}"
    elif "{" in rest or "}" in rest:
        why = "짝이 맞지 않는 중괄호"
    elif m := re.search(r"\d+", rest):
        why = f"자리표시자 밖의 숫자 '{m[0]}'"
    pct = anomaly_rules.change_pct(prev_v, new_v)
    vals = {"prev_value": f"{won(prev_v)}원", "new_value": f"{won(new_v)}원",
            "change_pct": "-" if pct is None else f"{pct:+.1f}%"}
    return fix_particles(PLACEHOLDER.sub(lambda x: vals[x[1]], DEFAULT_TEMPLATE if why else tmpl)), why


# ───────────────────────────── 자연어 조회: AI가 옮긴 표기를 코드가 해석
def resolve_nlq(plans: PlanBook, p: dict) -> tuple[dict, list[str]]:
    """요금제 표기 → 요금제 ID, 금액 표기 → 원, 가입유형 → 조건 키. (해석된 조건, 설명)"""
    notes, out = [], {k: p.get(k) for k in ("partner", "category", "field", "op", "direction")}
    out["plan_id"] = None
    if p.get("plan_phrase"):
        try:
            out["plan_id"] = plans.resolve(p.get("partner"), p["plan_phrase"]) or None
            notes.append(f"요금제 '{p['plan_phrase']}' → {out['plan_id']}")
        except UnknownCondition:
            notes.append(f"요금제 '{p['plan_phrase']}' → 사전에 없음(조건에서 뺌)")
    out["condition_key"] = JOIN_KEYS.get(p.get("join"))
    out["amount"] = parse_krw(p["amount_text"], gae=True) if p.get("amount_text") else None
    if p.get("amount_text"):
        notes.append(f"금액 '{p['amount_text']}' → {won(out['amount'])}원" if out["amount"] is not None
                     else f"금액 '{p['amount_text']}' → 해석 불가(조건에서 뺌)")
    return out, notes
